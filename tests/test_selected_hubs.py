"""Tests for the selected-hubs Excel report (summary + projects, Urgency flag, run date)."""

from datetime import datetime

import openpyxl
import pandas as pd
import pytest

from huburgency import run_end_to_end, selected_hubs_report
from huburgency.selected_hubs import (
    NOT_URGENT, URGENT, build_selected_hubs_tables, make_selected_hubs_report,
    read_hub_list, urgency_label,
)
from .test_inventar_formats import geojson_folder  # noqa: F401  (fixture)


# =============================================================================
# Pure pieces
# =============================================================================

def test_urgency_boundary_is_inclusive_at_50():
    pct = pd.Series([0.0, 49.99, 50.0, 50.01, 100.0])
    assert urgency_label(pct).tolist() == [NOT_URGENT, NOT_URGENT, URGENT, URGENT, URGENT]
    assert urgency_label(pct, threshold=80).tolist() == [NOT_URGENT] * 4 + [URGENT]


def _write(tmp_path, text, name='hubs.csv'):
    path = tmp_path / name
    path.write_text(text, encoding='utf-8')
    return path


def test_read_hub_list_variants(tmp_path):
    assert read_hub_list(_write(tmp_path, 'group\n3\n1\n3\n2\n')) == [3, 1, 2]     # order kept, deduped
    assert read_hub_list(_write(tmp_path, 'hub_id\n5\n6\n')) == [5, 6]             # single other column
    assert read_hub_list(_write(tmp_path, '5\n6\n7\n')) == [5, 6, 7]               # headerless
    assert read_hub_list(_write(tmp_path, 'group,name\n1,a\n2,b\n')) == [1, 2]     # extra columns ignored
    with pytest.raises(ValueError, match="group"):
        read_hub_list(_write(tmp_path, 'a,b\n1,2\n'))
    with pytest.raises(ValueError, match="No hub ids"):
        read_hub_list(_write(tmp_path, 'group\n'))


@pytest.fixture
def tables_input():
    progress = pd.DataFrame({
        'group': [1, 2, 3], 'HubNameHE': ['א', 'ב', 'ג'],
        'total_projects': [2, 0, 1], 'current_weighted_sum': [1.5, 0, 1.0],
        'max_possible_sum': [2.0, 0, 1.0], 'unique_statuses': [2, 0, 1],
        'status_progress_pct': [75.0, 0.0, 100.0],
    })
    joined = pd.DataFrame({
        'group': [1, 1, 1, 2, 3],
        'HubNameHE': ['א', 'א', 'א', 'ב', 'ג'],
        'project_uid': [10.0, 11.0, None, None, 12.0],    # NaN rows are zero-project placeholders
        'proj_name': ['p10', 'p11', None, None, 'p12'],
        'main_type': ['x'] * 5,
        'Proj_status': ['ביצוע', 'תכנון', None, None, 'ביצוע'],
        'Proj_status_original': ['ביצוע', '?', None, None, 'ביצוע'],
        'status_overridden': [False, True, None, None, False],
    })
    return progress, joined


def test_tables_filter_order_and_missing(tables_input):
    progress, joined = tables_input
    summary, projects, missing = build_selected_hubs_tables(progress, joined, [2, 1, 42])
    assert missing == [42]
    assert summary['group'].tolist() == [2, 1]                       # order of the list
    assert summary['Urgency'].tolist() == [NOT_URGENT, URGENT]
    assert list(summary.columns) == ['group', 'HubNameHE', 'total_projects', 'current_weighted_sum',
                                     'max_possible_sum', 'unique_statuses', 'status_progress_pct', 'Urgency']
    # hub 2: one empty row; hub 1: its two projects, the placeholder row is dropped; hub 3 not selected
    assert projects['group'].tolist() == [2, 1, 1]
    assert projects['project_uid'].tolist()[1:] == [10, 11]
    assert pd.isna(projects['project_uid'].iloc[0])
    assert projects['Urgency'].tolist() == [NOT_URGENT, URGENT, URGENT]
    assert projects.loc[projects['project_uid'] == 11, 'Proj_status_original'].item() == '?'


# =============================================================================
# Excel file
# =============================================================================

def _sheet_rows(path, name):
    ws = openpyxl.load_workbook(path)[name]
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def test_excel_has_run_date_summary_and_projects(tables_input, tmp_path):
    progress, joined = tables_input
    _write(tmp_path, 'group\n1\n2\n42\n', 'list.csv')
    report = make_selected_hubs_report(progress, joined, tmp_path / 'list.csv', tmp_path / 'r.xlsx',
                                       run_date=datetime(2026, 1, 2, 3, 4))
    summary = _sheet_rows(report.path, 'Summary')
    block = {r[0]: r[1] for r in summary if r[0] and r[1] is not None and r[2] is None}
    assert block['Run date'] == '2026-01-02 03:04'
    assert (block['Hubs requested'], block['Hubs in report'], block['Urgent hubs']) == (3, 2, 1)
    assert block['Hub ids not found'] == '42'
    assert block['Hub list'] == 'list.csv'
    header_row = next(i for i, r in enumerate(summary) if r[0] == 'group')
    assert summary[header_row][:8] == ['group', 'HubNameHE', 'total_projects', 'current_weighted_sum',
                                       'max_possible_sum', 'unique_statuses', 'status_progress_pct', 'Urgency']
    assert [r[0] for r in summary[header_row + 1:]] == [1, 2]
    assert [r[7] for r in summary[header_row + 1:]] == [URGENT, NOT_URGENT]

    projects = _sheet_rows(report.path, 'Projects')
    assert projects[0][:4] == ['group', 'HubNameHE', 'project_uid', 'proj_name']
    assert [(r[0], r[2]) for r in projects[1:]] == [(1, 10), (1, 11), (2, None)]


# =============================================================================
# Inside the workflow and standalone
# =============================================================================

def _selection(folder, ids):
    pd.DataFrame({'group': ids}).to_csv(folder / 'hub_selection.csv', index=False)


def test_workflow_writes_report_when_hub_list_exists(geojson_folder, tmp_path):   # noqa: F811
    _selection(geojson_folder, [2, 1, 99])
    out = tmp_path / 'out'
    result = run_end_to_end(geojson_folder, out, make_report=False)
    assert result.outputs['selected_hubs_report'] == out / 'selected_hubs_report.xlsx'
    assert result.selected_hubs.missing_groups == [99]
    summary = result.selected_hubs.summary_df
    assert summary['group'].tolist() == [2, 1]
    assert summary['status_progress_pct'].tolist() == [0.0, 75.0]     # hub 1: 'ביצוע' 1.0 + 'תכנון מפורט' .5
    assert summary['Urgency'].tolist() == [NOT_URGENT, URGENT]


def test_workflow_threshold_parameter(geojson_folder, tmp_path):                  # noqa: F811
    _selection(geojson_folder, [1])
    result = run_end_to_end(geojson_folder, tmp_path / 'out', make_report=False, urgent_threshold=80)
    assert result.selected_hubs.summary_df['Urgency'].tolist() == [NOT_URGENT]


def test_workflow_without_hub_list_writes_no_report(geojson_folder, tmp_path):    # noqa: F811
    out = tmp_path / 'out'
    result = run_end_to_end(geojson_folder, out, make_report=False)
    assert result.selected_hubs is None
    assert not (out / 'selected_hubs_report.xlsx').exists()


def test_link_only_writes_no_selected_report(geojson_folder, tmp_path):           # noqa: F811
    _selection(geojson_folder, [1])
    out = tmp_path / 'out'
    run_end_to_end(geojson_folder, out, link_only=True)
    assert not (out / 'selected_hubs_report.xlsx').exists()


def test_standalone_report_from_existing_results(geojson_folder, tmp_path):       # noqa: F811
    out = tmp_path / 'out'
    run_end_to_end(geojson_folder, out, make_report=False)
    list_csv = _write(tmp_path, 'group\n1\n', 'my_hubs.csv')
    report = selected_hubs_report(out, list_csv)
    assert report.path == out / 'selected_hubs_report.xlsx'
    assert report.summary_df['Urgency'].tolist() == [URGENT]
    projects = report.projects_df
    assert sorted(projects['project_uid'].tolist()) == [1, 2]
    assert projects.set_index('project_uid').loc[1, 'Proj_status'] == 'ביצוע'


def test_standalone_report_needs_existing_results(tmp_path):
    with pytest.raises(FileNotFoundError, match="run the full pipeline"):
        selected_hubs_report(tmp_path, _write(tmp_path, 'group\n1\n', 'l.csv'))
