"""
Selected-Hubs Report
====================

Excel report for a hand-picked list of hubs (a CSV of hub ids):

    Summary   run date, how many hubs, urgent count, ids not found, then one row
              per hub with the ``hub_progress`` columns and ``Urgency``
    Projects  one row per (hub, project) with the project's final status

``Urgency`` is "דחוף" when ``status_progress_pct >= urgent_threshold`` (default 50),
otherwise "לא דחוף".

Two ways to get it:
    * inside a full run: put ``hub_selection.csv`` in the input folder
      (``run_end_to_end(..., hub_selection_csv=...)``);
    * after a run, without re-linking::

        from huburgency.selected_hubs import selected_hubs_report
        selected_hubs_report('path/to/output', 'path/to/hub_list.csv')
"""

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Union
import logging

import pandas as pd

logger = logging.getLogger(__name__)

PathLike = Union[str, Path]

URGENT = 'דחוף'
NOT_URGENT = 'לא דחוף'
DEFAULT_THRESHOLD = 50.0
REPORT_FILENAME = 'selected_hubs_report.xlsx'

SUMMARY_COLUMNS = ['group', 'HubNameHE', 'total_projects', 'current_weighted_sum',
                   'max_possible_sum', 'unique_statuses', 'status_progress_pct', 'Urgency']
PROJECT_COLUMNS = ['group', 'HubNameHE', 'project_uid', 'proj_name', 'main_type', 'Proj_status',
                   'Proj_status_original', 'status_overridden', 'status_progress_pct', 'Urgency']


@dataclass
class SelectedHubsReport:
    """The report tables and where they were written."""

    path: Path
    summary_df: pd.DataFrame
    projects_df: pd.DataFrame
    requested: List[int]
    missing_groups: List[int]
    run_date: datetime
    urgent_threshold: float


# =============================================================================
# Reading the hub list
# =============================================================================

def read_hub_list(path: PathLike) -> List[int]:
    """
    Hub ids from a CSV, in file order without duplicates.

    Uses the ``group`` column; a single-column file with any other header (or no
    header at all) is accepted too.
    """
    from .workflow import read_csv_auto  # local import: workflow imports this module

    df = read_csv_auto(path)
    if 'group' in df.columns:
        values = df['group']
    elif len(df.columns) == 1:
        header = df.columns[0]
        values = df.iloc[:, 0]
        if pd.notna(pd.to_numeric(pd.Series([header]), errors='coerce').iloc[0]):
            values = pd.concat([pd.Series([header]), values], ignore_index=True)  # headerless list
    else:
        raise ValueError(
            f"{Path(path).name}: expected a 'group' column (or a single column of hub ids); "
            f"found {list(df.columns)}"
        )
    ids = pd.to_numeric(values, errors='coerce').dropna().astype('int64')
    ids = list(dict.fromkeys(ids.tolist()))
    if not ids:
        raise ValueError(f"No hub ids found in {Path(path).name}")
    return ids


# =============================================================================
# Building the tables
# =============================================================================

def urgency_label(progress_pct: pd.Series, threshold: float = DEFAULT_THRESHOLD) -> pd.Series:
    """'דחוף' where progress >= threshold, else 'לא דחוף'."""
    return progress_pct.ge(threshold).map({True: URGENT, False: NOT_URGENT})


def build_selected_hubs_tables(
    progress_df: pd.DataFrame,
    joined_df: pd.DataFrame,
    groups: List[int],
    urgent_threshold: float = DEFAULT_THRESHOLD,
    group_col: str = 'group',
):
    """Returns ``(summary_df, projects_df, missing_groups)`` for the requested hubs."""
    progress = progress_df.copy()
    progress[group_col] = pd.to_numeric(progress[group_col], errors='coerce')
    if 'HubNameHE' not in progress.columns:
        progress['HubNameHE'] = ''

    selected = progress[progress[group_col].isin(groups)].copy()
    missing = [g for g in groups if g not in set(selected[group_col].dropna().astype('int64'))]
    if missing:
        logger.warning(f"Hub ids not found in the results (left out of the report): {missing}")

    selected['Urgency'] = urgency_label(selected['status_progress_pct'], urgent_threshold)
    order = {g: i for i, g in enumerate(groups)}                # keep the order of the list
    selected = selected.assign(_o=selected[group_col].map(order)).sort_values('_o')
    summary = selected.rename(columns={group_col: 'group'})[SUMMARY_COLUMNS].reset_index(drop=True)
    summary['group'] = summary['group'].astype('int64')
    summary[['current_weighted_sum', 'max_possible_sum']] = \
        summary[['current_weighted_sum', 'max_possible_sum']].round(4)   # float noise (7.6499999...)

    joined = joined_df.copy()
    joined[group_col] = pd.to_numeric(joined[group_col], errors='coerce')
    joined = joined[joined[group_col].isin(groups)].rename(columns={group_col: 'group'})
    joined['project_uid'] = pd.to_numeric(joined['project_uid'], errors='coerce').astype('Int64')
    # Zero-project placeholder rows (NaN uid) only matter for hubs that have no project at all.
    has_project = joined.groupby('group')['project_uid'].transform(lambda s: s.notna().any())
    joined = joined[joined['project_uid'].notna() | ~has_project]
    joined = joined.drop_duplicates(subset=['group', 'project_uid'])
    for col in ('proj_name', 'main_type', 'Proj_status', 'Proj_status_original', 'status_overridden'):
        if col not in joined.columns:
            joined[col] = pd.NA
    hub_cols = summary[['group', 'HubNameHE', 'status_progress_pct', 'Urgency']]
    projects = (
        joined.drop(columns=[c for c in ('HubNameHE', 'status_progress_pct') if c in joined.columns])
        .merge(hub_cols, on='group', how='right')       # hubs without projects keep one empty row
        .assign(_o=lambda d: d['group'].map(order))
        .sort_values(['_o', 'project_uid'], na_position='first')
    )
    projects['group'] = projects['group'].astype('int64')
    projects = projects[PROJECT_COLUMNS].reset_index(drop=True)
    return summary, projects, missing


# =============================================================================
# Excel output
# =============================================================================

def write_selected_hubs_excel(report: SelectedHubsReport, hub_list_name: str = '') -> Path:
    """Write ``report`` to ``report.path`` (sheets 'Summary' and 'Projects')."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(bold=True, color='FFFFFF')
    urgent_fill = PatternFill('solid', fgColor='F8CBAD')

    summary, projects = report.summary_df, report.projects_df
    n_urgent = int((summary['Urgency'] == URGENT).sum())
    block = [
        ('Run date', report.run_date.strftime('%Y-%m-%d %H:%M')),
        ('Urgency rule', f"urgent ({URGENT}) if status_progress_pct >= {report.urgent_threshold:g}, "
                         f"otherwise not urgent ({NOT_URGENT})"),
        ('Hub list', hub_list_name or '-'),
        ('Hubs requested', len(report.requested)),
        ('Hubs in report', len(summary)),
        ('Urgent hubs', n_urgent),
        ('Not urgent hubs', len(summary) - n_urgent),
        ('Hub ids not found', ', '.join(map(str, report.missing_groups)) or 'none'),
    ]
    table_header_row = len(block) + 4          # title row, blank row, block, blank row

    report.path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(report.path, engine='openpyxl') as xw:
        summary.to_excel(xw, sheet_name='Summary', index=False, startrow=table_header_row - 1)
        projects = projects.astype(object).where(projects.notna(), None)   # empty cells, not <NA>
        projects.to_excel(xw, sheet_name='Projects', index=False)

        ws = xw.sheets['Summary']
        ws['A1'] = 'Selected hubs report'
        ws['A1'].font = Font(bold=True, size=14)
        for i, (label, value) in enumerate(block, start=3):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=value).alignment = Alignment(horizontal='left')

        for sheet, df, header_row in ((ws, summary, table_header_row),
                                      (xw.sheets['Projects'], projects, 1)):
            for cell in sheet[header_row]:
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
            sheet.auto_filter.ref = (f"A{header_row}:{get_column_letter(df.shape[1])}"
                                     f"{header_row + len(df)}")
            urgency_col = list(df.columns).index('Urgency') + 1
            pct_col = list(df.columns).index('status_progress_pct') + 1
            for r in range(header_row + 1, header_row + 1 + len(df)):
                sheet.cell(row=r, column=pct_col).number_format = '0.00'
                if sheet.cell(row=r, column=urgency_col).value == URGENT:
                    sheet.cell(row=r, column=urgency_col).fill = urgent_fill
            for c, name in enumerate(df.columns, start=1):
                widest = max([len(str(name))] + [len(str(v)) for v in df[name].head(500)
                                                 if pd.notna(v)])
                sheet.column_dimensions[get_column_letter(c)].width = min(max(14, widest + 2), 60)
        ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width or 0, 22)
    return report.path


def make_selected_hubs_report(
    progress_df: pd.DataFrame,
    joined_df: pd.DataFrame,
    hub_list_csv: PathLike,
    report_path: PathLike,
    urgent_threshold: float = DEFAULT_THRESHOLD,
    run_date: Optional[datetime] = None,
    group_col: str = 'group',
) -> SelectedHubsReport:
    """Build the tables for the hubs listed in ``hub_list_csv`` and write the Excel file."""
    groups = read_hub_list(hub_list_csv)
    summary, projects, missing = build_selected_hubs_tables(
        progress_df, joined_df, groups, urgent_threshold, group_col)
    report = SelectedHubsReport(
        path=Path(report_path), summary_df=summary, projects_df=projects, requested=groups,
        missing_groups=missing, run_date=run_date or datetime.now(),
        urgent_threshold=urgent_threshold,
    )
    write_selected_hubs_excel(report, Path(hub_list_csv).name)
    return report


def selected_hubs_report(
    output_dir: PathLike,
    hub_list_csv: PathLike,
    report_path: Optional[PathLike] = None,
    urgent_threshold: float = DEFAULT_THRESHOLD,
    run_date: Optional[datetime] = None,
) -> SelectedHubsReport:
    """
    Report for the hubs in ``hub_list_csv`` from the results already in ``output_dir``
    (``hub_progress.csv`` and ``hub_project_joined.csv``); nothing is re-linked or re-scored.
    """
    from .workflow import OUT_JOINED, OUT_PROGRESS, read_csv_auto

    output_dir = Path(output_dir)
    for name in (OUT_PROGRESS, OUT_JOINED):
        if not (output_dir / name).exists():
            raise FileNotFoundError(f"{output_dir / name} not found - run the full pipeline first.")
    return make_selected_hubs_report(
        read_csv_auto(output_dir / OUT_PROGRESS),
        read_csv_auto(output_dir / OUT_JOINED),
        hub_list_csv,
        report_path or output_dir / REPORT_FILENAME,
        urgent_threshold=urgent_threshold,
        run_date=run_date,
    )
